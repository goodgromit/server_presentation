# -*- coding: utf-8 -*-
"""발표대본.md를 읽어 발표슬라이드.pptx를 만든다. 슬라이드 노트에는 대본이 들어간다."""
import re
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).parent
FONT = "맑은 고딕"
MONO = "Consolas"

INK = "1F2937"
MUTED = "6B7280"
LINE = "D1D5DB"
PANEL = "F8FAFC"
NAVY = "0F172A"
ACCENT = "2563EB"
WARN = "C2410C"
WARN_T = "FFF7ED"
GOOD = "047857"
WHITE = "FFFFFF"

# 구성 요소별 색: 모든 슬라이드에서 같은 의미로 쓴다
CLIENT, CLIENT_T = "2563EB", "DBEAFE"
API, API_T = "0F766E", "CCFBF1"
MATCH, MATCH_T = "6D28D9", "EDE9FE"
ROOM, ROOM_T = "C2410C", "FFEDD5"
DB, DB_T = "475569", "E2E8F0"
CACHE, CACHE_T = "0E7490", "CFFAFE"

ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
TOTAL = 22


def rgb(h):
    return RGBColor.from_string(h)


def strip_style(shape):
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def set_run(run, size, bold, color, font):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", font)


def fill_tf(tf, paras, size=16, bold=False, color=INK, align="l", font=FONT, after=0, spacing=None):
    """paras: 문자열(\n으로 문단 구분) 또는 문단 목록.
    문단은 문자열, (text, opts) 런 목록, 또는 {'runs':..., 'size':..., ...} 사전."""
    if isinstance(paras, str):
        paras = paras.split("\n")
    first = True
    for p in paras:
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        d = dict(size=size, bold=bold, color=color, font=font, align=align, after=after, spacing=spacing)
        if isinstance(p, dict):
            d.update({k: v for k, v in p.items() if k != "runs"})
            runs = p["runs"]
        else:
            runs = p
        if isinstance(runs, str):
            runs = [(runs, {})]
        para.alignment = ALIGN[d["align"]]
        para.space_after = Pt(d["after"])
        if d["spacing"]:
            para.line_spacing = d["spacing"]
        for t, o in runs:
            r = para.add_run()
            r.text = t
            ro = {**d, **o}
            set_run(r, ro["size"], ro["bold"], ro["color"], ro["font"])


def tx(s, x, y, w, h, paras, anchor="t", margin=0.04, **kw):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = ANCHOR[anchor]
    fill_tf(tf, paras, **kw)
    return tb


def box(s, x, y, w, h, paras=None, fill=None, line=None, lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
        radius=0.12, anchor="m", align="c", margin=0.1, dash=False, **kw):
    sp = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    strip_style(sp)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sp.adjustments[0] = radius
    if fill:
        sp.fill.solid()
        sp.fill.fore_color.rgb = rgb(fill)
    else:
        sp.fill.background()
    if line:
        sp.line.color.rgb = rgb(line)
        sp.line.width = Pt(lw)
        if dash:
            ln = sp.line._get_or_add_ln()
            etree.SubElement(ln, qn("a:prstDash")).set("val", "dash")
    else:
        sp.line.fill.background()
    tf = sp.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    tf.vertical_anchor = ANCHOR[anchor]
    if paras is not None:
        fill_tf(tf, paras, align=align, **kw)
    return sp


def _line_ends(ln, color, w, head, tail, dash):
    ln.set("w", str(Pt(w)))
    sf = etree.SubElement(ln, qn("a:solidFill"))
    etree.SubElement(sf, qn("a:srgbClr")).set("val", color)
    if dash:
        etree.SubElement(ln, qn("a:prstDash")).set("val", "dash")
    etree.SubElement(ln, qn("a:round"))
    if tail:
        e = etree.SubElement(ln, qn("a:headEnd"))
        e.set("type", "triangle"); e.set("w", "med"); e.set("len", "med")
    if head:
        e = etree.SubElement(ln, qn("a:tailEnd"))
        e.set("type", "triangle"); e.set("w", "med"); e.set("len", "med")


def arrow(s, x1, y1, x2, y2, color=MUTED, w=1.75, head=True, tail=False, dash=False):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    strip_style(c)
    spPr = c._element.spPr
    ln = spPr.find(qn("a:ln"))
    if ln is not None:
        spPr.remove(ln)
    _line_ends(etree.SubElement(spPr, qn("a:ln")), color, w, head, tail, dash)
    return c


def path(s, pts, color=MUTED, w=1.75, head=True, dash=False):
    """꺾인 화살표. pts는 인치 좌표 목록."""
    fb = s.shapes.build_freeform(Inches(pts[0][0]), Inches(pts[0][1]), scale=1.0)
    fb.add_line_segments([(Inches(x), Inches(y)) for x, y in pts[1:]], close=False)
    sp = fb.convert_to_shape()
    strip_style(sp)
    sp.fill.background()
    spPr = sp._element.spPr
    ln = spPr.find(qn("a:ln"))
    if ln is not None:
        spPr.remove(ln)
    _line_ends(etree.SubElement(spPr, qn("a:ln")), color, w, head, False, dash)
    return sp


def table(s, x, y, widths, rows, row_h=0.55, head_h=0.5, size=14, head_fill=INK, bold_first=True,
          zebra=("FFFFFF", "F3F4F6"), col_colors=None):
    shp = s.shapes.add_table(len(rows), len(widths), Inches(x), Inches(y), Inches(sum(widths)),
                             Inches(head_h + row_h * (len(rows) - 1)))
    tbl = shp.table
    tbl.first_row = True
    tbl.horz_banding = False
    for i, w in enumerate(widths):
        tbl.columns[i].width = Inches(w)
    for r, row in enumerate(rows):
        tbl.rows[r].height = Inches(head_h if r == 0 else row_h)
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = rgb(head_fill)
                color, bold = WHITE, True
            else:
                cell.fill.fore_color.rgb = rgb(zebra[(r - 1) % 2])
                color = (col_colors or {}).get(c, INK)
                bold = bold_first and c == 0
            cell.margin_left = cell.margin_right = Inches(0.14)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            tf.paragraphs[0].text = ""
            if isinstance(val, list) and val and isinstance(val[0], tuple):
                val = [val]  # 셀 하나 = 런 목록 한 문단
            fill_tf(tf, val, size=size - (1 if r == 0 else 0), bold=bold, color=color)
    return shp


def tag(s, label, color=ACCENT, fill="EFF6FF", x=10.75, w=1.98):
    box(s, x, 0.4, w, 0.36, label, fill=fill, line=color, lw=0.75, size=11, bold=True, color=color, radius=0.5)


def warn_icon(s, x, y, d=0.42):
    box(s, x, y, d, d, "!", fill=WARN, shape=MSO_SHAPE.OVAL, size=16, bold=True, color=WHITE, margin=0)


def content_slide(part, title):
    s = prs.slides.add_slide(BLANK)
    n = len(prs.slides)
    tx(s, 0.6, 0.38, 9.5, 0.32, part, size=12, bold=True, color=ACCENT)
    tx(s, 0.6, 0.66, 12.1, 0.75, title, size=28, bold=True, color=INK)
    tx(s, 0.6, 7.02, 7, 0.3, "Level 9 · 클라이언트 서버 통신과 멀티플레이", size=10, color=MUTED)
    tx(s, 11.4, 7.02, 1.33, 0.3, f"{n} / {TOTAL}", size=10, color=MUTED, align="r")
    return s


def dark_slide():
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(NAVY)
    return s


P1 = "PART 1 · 게임 흐름과 책임"
P2 = "PART 2 · 서버 구성 요소"
P3 = "PART 3 · 통신 방식"
P4 = "PART 4 · 데이터 · 매칭 · 실시간"
P5 = "PART 5 · 데모와 AI 활용"

# ---------------------------------------------------------------- 1. 표지
s = dark_slide()
box(s, 0.9, 1.55, 0.12, 2.3, fill=ACCENT, shape=MSO_SHAPE.RECTANGLE)
tx(s, 1.25, 1.45, 10, 0.5, "LEVEL 9", size=18, bold=True, color="93C5FD")
tx(s, 1.25, 1.95, 11, 1.3, "클라이언트 서버 통신과 멀티플레이", size=44, bold=True, color=WHITE)
tx(s, 1.25, 3.15, 11, 0.7, "소규모(2~5인) 협동·경쟁 게임으로 보는 서버와 데이터의 흐름", size=20, color="CBD5E1")
tx(s, 1.25, 5.6, 11, 0.4, "함께 보는 자료 · 서버 타입과 통신 방식 핸드아웃", size=14, color="94A3B8")

# ---------------------------------------------------------------- 2. 진행
s = content_slide("오늘의 진행", "SDK 사용법보다 구조를 본다")
agenda = [
    ("01", "게임 흐름과 책임", "게임 한 판의 흐름 · 누가 무엇을 책임지는가", "6분"),
    ("02", "서버 구성 요소", "API · 매칭 · 게임 룸 · 판정 주체 · 저장소", "24분"),
    ("03", "통신 방식", "TCP · UDP · HTTP REST · WebSocket", "12분"),
    ("04", "데이터 · 매칭 · 실시간", "Firebase 저장 · 매칭 상태도 · 예외 · 최적화", "16분"),
    ("05", "데모와 AI 활용", "Firebase 매칭 데모 · AI 초안 검토와 검증", "12분"),
    ("—", "정리와 질의응답", "설계 순서 6단계 · 여러분의 프로젝트", "15분"),
]
for i, (num, title, topics, mins) in enumerate(agenda):
    y = 1.65 + i * 0.78
    tx(s, 0.6, y, 0.9, 0.6, num, size=24, bold=True, color=ACCENT, anchor="m")
    tx(s, 1.55, y, 4.1, 0.6, title, size=20, bold=True, anchor="m")
    tx(s, 5.7, y, 5.6, 0.6, topics, size=14, color=MUTED, anchor="m")
    tx(s, 11.4, y, 1.33, 0.6, mins, size=14, color=MUTED, align="r", anchor="m")
    if i < len(agenda) - 1:
        arrow(s, 0.6, y + 0.7, 12.73, y + 0.7, color="E5E7EB", w=0.75, head=False)
tx(s, 0.6, 6.45, 12.1, 0.4,
   [[("세 가지 질문을 계속 묻는다  ", {"bold": True, "color": INK}),
     ("누가 결정하는가 · 무엇을 어디에 저장하는가 · 어떻게 전달하는가", {"color": ACCENT, "bold": True})]],
   size=15)

# ---------------------------------------------------------------- 3. 게임 흐름
s = content_slide(P1, "오늘의 게임 흐름")
steps = [("로그인", "API 서버", API, API_T), ("프로필 로드", "API · DB", API, API_T),
         ("로비", "API · 알림", API, API_T), ("매칭", "매칭 서버", MATCH, MATCH_T),
         ("게임 룸", "실시간 룸", ROOM, ROOM_T), ("결과 저장", "API · DB", API, API_T)]
bw, gap, x0 = 1.75, 0.31, 0.6
for i, (name, srv, c, t) in enumerate(steps):
    x = x0 + i * (bw + gap)
    box(s, x, 1.75, bw, 0.95, name, fill=t, line=c, lw=1.25, size=17, bold=True)
    tx(s, x, 2.75, bw, 0.35, srv, size=12, bold=True, color=c, align="c")
    if i < len(steps) - 1:
        arrow(s, x + bw + 0.04, 2.225, x + bw + gap - 0.04, 2.225, color=MUTED, w=1.5)
tx(s, 0.6, 3.4, 12, 0.4, "같은 흐름, 세 가지 시선", size=16, bold=True)
views = [("기획자", "규칙으로 본다", "몇 명이 모이면 시작하나?\n중간에 나가면 패배인가?"),
         ("클라이언트 개발자", "화면과 입력으로 본다", "매칭 중에는 어떤 화면을 띄우나?\n취소 버튼은 언제 막나?"),
         ("서버 개발자", "권한 · 데이터 · 장애로 본다", "이 요청은 누가 보낼 수 있나?\n실패하면 어떻게 되돌리나?")]
for i, (who, how, q) in enumerate(views):
    x = 0.6 + i * 4.14
    box(s, x, 3.9, 3.85, 2.45, fill=PANEL, line="E5E7EB", radius=0.06)
    tx(s, x + 0.25, 4.05, 3.4, 0.4, who, size=13, bold=True, color=ACCENT)
    tx(s, x + 0.25, 4.42, 3.4, 0.5, how, size=19, bold=True)
    tx(s, x + 0.25, 5.05, 3.4, 1.2, q, size=14, color="374151", after=4)

# ---------------------------------------------------------------- 4. 책임
s = content_slide(P1, "누가 무엇을 책임지는가")
cards = [("클라이언트", "빠르게 보여 준다", ["화면", "입력", "로컬 표현"], CLIENT),
         ("서버", "판단하고 확정한다", ["요청 처리", "상태 관리", "권한 검증"], API),
         ("DB · Firebase", "기억한다", ["다시 접속해도", "남는 데이터"], DB)]
for i, (name, verb, items, c) in enumerate(cards):
    x = 0.6 + i * 4.14
    box(s, x, 1.7, 3.85, 2.65, fill=PANEL, line="E5E7EB", radius=0.05)
    box(s, x, 1.7, 3.85, 0.1, fill=c, shape=MSO_SHAPE.RECTANGLE)
    tx(s, x + 0.3, 1.95, 3.3, 0.5, name, size=22, bold=True, color=c)
    tx(s, x + 0.3, 2.45, 3.3, 0.4, verb, size=14, color=MUTED)
    tx(s, x + 0.3, 2.95, 3.3, 1.3, "\n".join(items), size=18, bold=True, after=2)
box(s, 0.6, 4.65, 5.65, 1.95, fill=WARN_T, line="FED7AA", radius=0.06)
tx(s, 0.85, 4.8, 5.2, 0.45, "보여 주는 값 ≠ 확정된 값", size=18, bold=True, color=WARN)
tx(s, 0.85, 5.3, 5.2, 1.2,
   ["내 화면: 총알이 적에게 명중!",
    "서버 판정: 그 순간 적은 이미 엄폐물 뒤",
    [("→ 데미지는 들어가지 않는다", {"bold": True})]], size=14, after=3)
box(s, 6.5, 4.65, 6.23, 1.95, fill=NAVY, radius=0.06)
tx(s, 6.8, 4.85, 5.7, 0.7, "누가 결과를 확정하는가?", size=28, bold=True, color=WHITE)
tx(s, 6.8, 5.6, 5.7, 0.8, "서버 구조를 결정하는 가장 중요한 질문\n“나 이겼어, 보상 줘”를 그대로 믿어도 될까?", size=14, color="CBD5E1", after=2)

# ---------------------------------------------------------------- 5. 서버 타입
s = content_slide(P2, "서버 타입 한눈에 보기")
tag(s, "핸드아웃 §1 · §2")
box(s, 0.6, 3.3, 1.6, 1.0, "클라이언트", fill=CLIENT_T, line=CLIENT, lw=1.25, size=15, bold=True)
box(s, 3.2, 1.75, 2.0, 0.85, "API 서버", fill=API_T, line=API, lw=1.25, size=15, bold=True)
box(s, 3.2, 3.35, 2.0, 0.85, "매칭 서버", fill=MATCH_T, line=MATCH, lw=1.25, size=15, bold=True)
box(s, 3.2, 4.95, 2.0, 0.85, [{"runs": "실시간 게임 룸", "bold": True},
                              {"runs": "릴레이 또는 전용 서버", "size": 11, "color": MUTED}],
    fill=ROOM_T, line=ROOM, lw=1.25, size=15)
box(s, 5.95, 1.65, 1.2, 1.05, "Firebase\n· DB", fill=DB_T, line=DB, lw=1.25, shape=MSO_SHAPE.CAN, size=13, bold=True)
arrow(s, 2.2, 3.55, 3.2, 2.3, color=API)
tx(s, 0.8, 2.95, 1.55, 0.3, "HTTP/REST", size=11, bold=True, color=API, align="r")
arrow(s, 2.2, 3.78, 3.2, 3.78, color=MATCH, tail=True)
tx(s, 2.15, 3.42, 1.1, 0.3, "매칭 요청", size=11, bold=True, color=MATCH, align="c")
arrow(s, 2.2, 4.05, 3.2, 5.25, color=ROOM, tail=True)
tx(s, 0.8, 4.45, 1.65, 0.3, "실시간 통신", size=11, bold=True, color=ROOM, align="r")
arrow(s, 5.2, 2.17, 5.95, 2.17, color=DB, tail=True)
path(s, [(5.2, 5.2), (5.6, 5.2), (5.6, 2.45), (5.2, 2.45)], color=ROOM)
tx(s, 5.68, 3.55, 1.3, 0.6, "검증된\n결과", size=11, bold=True, color=ROOM)
table(s, 7.45, 1.75, [1.85, 3.43], [
    ["구성 요소", "한 줄 역할"],
    [[("API 서버", {"color": API})], "요청을 받아 처리하고 응답한다"],
    [[("매칭 서버", {"color": MATCH})], "대기열을 관리하고 게임 룸을 배정한다"],
    [[("실시간 게임 룸", {"color": ROOM})], "게임 중 이벤트와 상태를 교환한다"],
    [[("전용 게임 서버", {"color": ROOM})], "게임 상태 · 승패를 권한 있게 확정"],
], row_h=0.78, size=13)
box(s, 0.6, 6.05, 12.13, 0.7,
    [[("작은 프로젝트에서는 여러 역할이 하나의 서비스에 합쳐질 수 있다.  ", {}),
      ("→ 이름보다 책임과 데이터 흐름을 먼저 나눈다", {"bold": True, "color": ACCENT})]],
    fill="EFF6FF", size=15, align="l", margin=0.25)

# ---------------------------------------------------------------- 6. API 서버
s = content_slide(P2, "API 서버는 요청과 응답을 처리한다")
reqs = [("POST /login", "인증 · 프로필 확인"), ("GET /profile", "프로필 반환"), ("POST /result", "결과 검증 · 저장")]
for i, (ep, res) in enumerate(reqs):
    y = 1.7 + i * 0.68
    box(s, 0.6, y, 1.55, 0.52, "클라이언트", fill=CLIENT_T, line=CLIENT, size=13, bold=True)
    arrow(s, 2.2, y + 0.26, 2.55, y + 0.26)
    box(s, 2.6, y, 2.4, 0.52, ep, fill="1E293B", size=15, bold=True, color=WHITE, font=MONO)
    arrow(s, 5.05, y + 0.26, 5.4, y + 0.26)
    box(s, 5.45, y, 1.55, 0.52, "API 서버", fill=API_T, line=API, size=13, bold=True)
    arrow(s, 7.05, y + 0.26, 7.4, y + 0.26)
    tx(s, 7.5, y, 3.0, 0.52, res, size=15, bold=True, anchor="m")
tx(s, 10.3, 1.7, 2.43, 1.9, "로그인 · 프로필\n인벤토리 · 상점\n전적 · 보상\n\n즉시 답을 받는 업무에 적합", size=13, color=MUTED)
box(s, 0.6, 3.95, 5.95, 2.75, fill=WARN_T, line="FED7AA", radius=0.05)
warn_icon(s, 0.85, 4.15)
tx(s, 1.4, 4.12, 5.0, 0.5, "두 번 눌러도 한 번만 처리", size=19, bold=True, color=WARN)
tx(s, 0.85, 4.75, 5.5, 1.9, [
    [("상황  ", {"bold": True}), ("응답을 못 받은 클라이언트가 보상 요청을 다시 보냄", {})],
    [("위험  ", {"bold": True}), ("첫 요청은 이미 처리됨 → 보상 2번 지급", {})],
    [("해결  ", {"bold": True}), ("요청 ID로 처리 기록을 남기고 첫 결과를 다시 반환 (멱등성)", {})],
], size=14, after=6)
tx(s, 6.85, 3.95, 5.9, 0.45, "API 문서에 함께 적을 4가지", size=17, bold=True)
docs = [("요청 데이터", "matchId, score, 사용자 토큰"), ("성공 응답", "200 + 지급 내역"),
        ("실패 응답", "이미 처리됨 · 참가자 아님 · 없는 경기"), ("권한 조건", "로그인한 경기 참가자만")]
for i, (k, v) in enumerate(docs):
    x = 6.85 + (i % 2) * 2.98
    y = 4.45 + (i // 2) * 1.15
    box(s, x, y, 2.88, 1.05, fill=PANEL, line="E5E7EB", radius=0.08)
    tx(s, x + 0.15, y + 0.1, 2.6, 0.35, k, size=15, bold=True, color=API)
    tx(s, x + 0.15, y + 0.47, 2.6, 0.55, v, size=12, color="374151")

# ---------------------------------------------------------------- 7. 매칭 vs 룸
s = content_slide(P2, "매칭 서버와 게임 룸은 다르다")
lanes = [("매칭 서버", "자리를 안내하는 직원", ["대기열 등록", "조건 비교", "방 ID 배정"],
          "목표: 같은 조건의 사용자를 빠르게 묶어 방에 넣는다", MATCH, MATCH_T),
         ("게임 룸", "손님이 앉는 테이블", ["방 입장", "준비", "실제 게임 데이터 교환"],
          "목표: 매칭된 사용자들이 끝까지 플레이하도록 연결을 유지한다", ROOM, ROOM_T)]
for li, (name, meta, sts, goal, c, t) in enumerate(lanes):
    y = 1.7 + li * 1.8
    box(s, 0.6, y, 2.6, 1.55, [{"runs": name, "size": 20, "bold": True, "color": WHITE},
                               {"runs": meta, "size": 12, "color": WHITE}], fill=c, radius=0.08)
    for i, st in enumerate(sts):
        x = 3.55 + i * 3.1
        box(s, x, y + 0.12, 2.6, 0.75, st, fill=t, line=c, lw=1.25, size=16, bold=True)
        if i < 2:
            arrow(s, x + 2.65, y + 0.495, x + 3.05, y + 0.495, color=c)
    tx(s, 3.55, y + 1.0, 8.9, 0.4, goal, size=14, color="374151")
tx(s, 10.3, 2.82, 2.4, 0.3, "방 ID 전달 ↓", size=12, bold=True, color=MATCH, align="c")
box(s, 0.6, 5.45, 12.13, 1.3, fill=WARN_T, line="FED7AA", radius=0.06)
warn_icon(s, 0.85, 5.65)
tx(s, 1.4, 5.6, 11.1, 0.45, "취소와 매칭 확정이 동시에 일어나면?", size=18, bold=True, color=WARN)
tx(s, 1.4, 6.08, 11.1, 0.6, "한쪽은 ‘취소됨’, 상대는 ‘매칭 완료’ → 최종 상태를 미리 정하고, 서버는 상태 변경을 하나씩 처리한다",
   size=14)

# ---------------------------------------------------------------- 8. 릴레이 vs 전용
s = content_slide(P2, "실시간 릴레이와 전용 게임 서버")
cols = [("실시간 릴레이 룸", "우체국: 메시지를 전달만 한다", "64748B",
         [("판정 주체", "클라이언트 또는 호스트"), ("적합한 상황", "캐주얼 협동, 빠른 프로토타입"),
          ("클라이언트가 보내는 것", "상태와 결과")]),
        ("전용 게임 서버", "심판: 게임 루프를 돌리고 확정한다", "1E3A8A",
         [("판정 주체", "서버"), ("적합한 상황", "경쟁 PvP, 랭킹, 치팅 방지"),
          ("클라이언트가 보내는 것", "입력만 (“공격 버튼 눌렀음”)")])]
for i, (name, meta, c, rows) in enumerate(cols):
    x = 0.6 + i * 6.18
    box(s, x, 1.7, 5.95, 3.4, fill=PANEL, line="E5E7EB", radius=0.04)
    box(s, x, 1.7, 5.95, 0.95, [{"runs": name, "size": 21, "bold": True, "color": WHITE},
                                {"runs": meta, "size": 13, "color": "E2E8F0"}],
        fill=c, shape=MSO_SHAPE.RECTANGLE, align="l", margin=0.3)
    for j, (k, v) in enumerate(rows):
        y = 2.85 + j * 0.74
        tx(s, x + 0.3, y, 5.4, 0.3, k, size=12, color=MUTED)
        tx(s, x + 0.3, y + 0.27, 5.4, 0.42, v, size=17, bold=True)
arrow(s, 0.9, 5.55, 12.4, 5.55, color="94A3B8", w=2.5)
tx(s, 0.9, 5.15, 4, 0.35, "릴레이", size=12, bold=True, color=MUTED)
tx(s, 8.4, 5.15, 4, 0.35, "전용 서버", size=12, bold=True, color=MUTED, align="r")
box(s, 2.6, 5.75, 3.6, 0.5, "서버 비용 · 구현 난도  ↑", fill="F1F5F9", size=14, bold=True, color="475569", radius=0.5)
box(s, 7.1, 5.75, 3.6, 0.5, "공정성 · 검증 능력  ↑", fill="DBEAFE", size=14, bold=True, color="1E3A8A", radius=0.5)
tx(s, 0.6, 6.42, 12.13, 0.45, "메시지가 서버를 지나가도, 서버가 게임 규칙을 판정하는 것은 아니다.", size=16, bold=True, color=WARN, align="c")

# ---------------------------------------------------------------- 9. 판정 구조 4가지
s = content_slide(P2, "누가 게임 결과를 확정하는가: 네 가지 구조")
tag(s, "핸드아웃 §5")
arrow(s, 1.0, 1.82, 12.3, 1.82, color="94A3B8", w=2.25)
tx(s, 0.6, 1.5, 3, 0.3, "빠르게 만들기 쉬움", size=12, bold=True, color=MUTED)
tx(s, 9.7, 1.5, 3, 0.3, "공정한 판정에 유리", size=12, bold=True, color=MUTED, align="r")
models = [("클라이언트 주도", "각 플레이어 클라이언트", "빠르게 만들기 쉽다", "조작과 불일치에 약하다"),
          ("호스트 클라이언트", "방장 역할의 클라이언트", "별도 서버 비용이 적다", "호스트 이탈 · 호스트 조작"),
          ("릴레이 룸", "클라이언트들 (공유 상태)", "소규모 협동에 단순하다", "권한 있는 판정 서버 아님"),
          ("전용 게임 서버", "서버", "공정한 판정과 검증", "서버 호스팅과 운영 필요")]
for i, (name, who, pro, con) in enumerate(models):
    x = 0.6 + i * 3.08
    hl = i == 3
    box(s, x, 2.1, 2.9, 3.15, fill="EFF6FF" if hl else PANEL, line=ACCENT if hl else "E5E7EB",
        lw=1.5 if hl else 1, radius=0.05)
    tx(s, x + 0.2, 2.22, 2.5, 0.3, f"0{i + 1}", size=13, bold=True, color=ACCENT)
    tx(s, x + 0.2, 2.5, 2.6, 0.45, name, size=18, bold=True)
    tx(s, x + 0.2, 2.95, 2.6, 0.55, [[("판정  ", {"color": MUTED}), (who, {"bold": True})]], size=12)
    tx(s, x + 0.2, 3.55, 2.6, 0.3, "장점", size=12, bold=True, color=GOOD)
    tx(s, x + 0.2, 3.82, 2.6, 0.55, pro, size=14)
    tx(s, x + 0.2, 4.4, 2.6, 0.3, "위험 · 비용", size=12, bold=True, color=WARN)
    tx(s, x + 0.2, 4.67, 2.6, 0.55, con, size=14)
box(s, 0.6, 5.45, 12.13, 1.3, fill=NAVY, radius=0.06)
tx(s, 0.95, 5.58, 11.5, 0.5, "전적 · 보상 · 랭킹은 신뢰할 수 있는 서버 검증을 거친 뒤 저장한다", size=20, bold=True, color=WHITE)
tx(s, 0.95, 6.12, 11.5, 0.5, "플레이는 릴레이로 가볍게, 저장하는 순간에는 서버가 검증  (예: 5분 경기를 10초 만에 이겼다는 결과는 거부)",
   size=14, color="CBD5E1")

# ---------------------------------------------------------------- 10. 저장소
s = content_slide(P2, "DB, 캐시, Firebase의 역할")
stores = [("수명: 몇 달 이상", "DB · Firebase", "영구 · 정확해야 한다", ["계정 · 프로필", "전적 · 보상", "구매 기록"], DB),
          ("수명: 몇 분 ~ 몇 시간", "캐시", "빠르지만 사라질 수 있다", ["세션", "대기열", "짧게 유지할 룸 상태"], CACHE),
          ("수명: 몇 초", "실시간 룸", "곧 의미가 없어진다", ["입력 · 위치", "준비 상태", "순간 이벤트"], ROOM)]
for i, (life, name, prop, items, c) in enumerate(stores):
    x = 0.6 + i * 4.14
    box(s, x, 1.7, 3.85, 3.3, fill=PANEL, line="E5E7EB", radius=0.05)
    box(s, x, 1.7, 3.85, 0.1, fill=c, shape=MSO_SHAPE.RECTANGLE)
    tx(s, x + 0.3, 1.95, 3.3, 0.3, life, size=12, bold=True, color=c)
    tx(s, x + 0.3, 2.25, 3.3, 0.5, name, size=22, bold=True)
    tx(s, x + 0.3, 2.78, 3.3, 0.35, prop, size=13, color=MUTED)
    tx(s, x + 0.3, 3.3, 3.3, 1.6, ["•  " + t for t in items], size=16, after=4)
for i, (h, b) in enumerate([("실시간 위치를 모두 영구 DB에 쓰면", "5명 × 초당 20회 × 10분 = 60,000번 쓰기\n→ 비용과 지연이 커진다"),
                            ("보상 지급 기록을 캐시에만 두면", "캐시가 사라지면 기록도 사라진다\n→ 중복 지급 또는 누락")]):
    x = 0.6 + i * 6.18
    box(s, x, 5.25, 5.95, 1.5, fill=WARN_T, line="FED7AA", radius=0.06)
    warn_icon(s, x + 0.25, 5.45)
    tx(s, x + 0.8, 5.4, 5.0, 0.45, h, size=16, bold=True, color=WARN)
    tx(s, x + 0.8, 5.85, 5.0, 0.8, b, size=14)

# ---------------------------------------------------------------- 11. TCP
s = content_slide(P3, "연결지향 통신: TCP")
tag(s, "핸드아웃 §3")
tx(s, 0.6, 1.7, 5.4, 1.1, "연결을 만들고 유지하며,\n순서와 전달을 보장한다.", size=24, bold=True, spacing=1.1)
tx(s, 0.6, 2.95, 5.4, 0.35, "비유: 전화 통화 · 연결부터 맺고 이야기한다", size=13, color=MUTED)
tx(s, 0.6, 3.5, 5.4, 0.35, "적합한 데이터", size=14, bold=True, color=ACCENT)
for i, t in enumerate(["로그인", "프로필 · 인벤토리", "구매 · 보상", "경기 결과 저장"]):
    box(s, 0.6 + (i % 2) * 2.7, 3.9 + (i // 2) * 0.68, 2.55, 0.55, t, fill="EFF6FF", size=15, bold=True, color="1E3A8A")
box(s, 6.2, 1.7, 6.53, 3.95, fill=PANEL, line="E5E7EB", radius=0.04)
tx(s, 6.4, 1.8, 6.2, 0.35, "패킷 2가 사라지면", size=14, bold=True)
labels = [("보냄", 2.35), ("도착", 3.25), ("앱에 전달", 4.45)]
for lab, y in labels:
    tx(s, 6.4, y, 1.4, 0.55, lab, size=13, bold=True, color=MUTED, anchor="m")
px = [8.0, 8.7, 9.4, 10.1]
for i, x in enumerate(px):
    box(s, x, 2.35, 0.55, 0.55, str(i + 1), fill=CLIENT_T, line=CLIENT, size=16, bold=True, radius=0.15)
for i, x in enumerate(px):
    if i == 1:
        box(s, x, 3.25, 0.55, 0.55, "✕", line="DC2626", dash=True, size=16, bold=True, color="DC2626", radius=0.15)
    else:
        box(s, x, 3.25, 0.55, 0.55, str(i + 1), fill=CLIENT_T, line=CLIENT, size=16, bold=True, radius=0.15)
box(s, 11.3, 3.25, 0.55, 0.55, "2", fill=ROOM_T, line=ROOM, size=16, bold=True, radius=0.15)
tx(s, 10.85, 3.82, 1.6, 0.3, "재전송", size=11, bold=True, color=ROOM, align="c")
box(s, 8.0, 4.45, 0.55, 0.55, "1", fill=CLIENT_T, line=CLIENT, size=16, bold=True, radius=0.15)
box(s, 8.7, 4.45, 2.15, 0.55, "3 · 4는 도착했지만 대기", line="94A3B8", dash=True, size=11, color=MUTED, radius=0.15)
for i, x in enumerate([10.95, 11.5, 12.05]):
    box(s, x, 4.45, 0.5, 0.55, str(i + 2), fill=CLIENT_T, line=CLIENT, size=16, bold=True, radius=0.15)
tx(s, 6.4, 5.12, 6.2, 0.4, "손실은 재전송하고, 앱에는 항상 1 → 2 → 3 → 4 순서로 넘긴다", size=12, color=MUTED)
box(s, 0.6, 5.95, 12.13, 0.8,
    [[("앞 데이터가 늦으면 뒤 데이터도 기다린다.  ", {"bold": True, "color": WARN}),
      ("빠르게 변하는 위치값에는 불리할 수 있다 (턴제·느린 협동이라면 TCP로 충분)", {})]],
    fill=WARN_T, size=14, align="l", margin=0.3)

# ---------------------------------------------------------------- 12. UDP
s = content_slide(P3, "비연결지향 통신: UDP")
tag(s, "핸드아웃 §3")
tx(s, 0.6, 1.7, 5.9, 1.1, "각 패킷은 독립적으로 전송된다.\n최신 데이터가 더 중요할 때 쓴다.", size=22, bold=True, spacing=1.1)
for i, t in enumerate(["이동 입력", "위치 스냅샷", "빠른 액션 게임 상태"]):
    w = [1.5, 1.7, 2.45][i]
    x = 0.6 + [0, 1.6, 3.4][i]
    box(s, x, 2.95, w, 0.5, t, fill=ROOM_T, size=14, bold=True, color=ROOM, radius=0.5)
box(s, 0.6, 3.7, 5.9, 2.6, fill=PANEL, line="E5E7EB", radius=0.04)
for lab, y in [("보냄", 3.95), ("도착", 4.75)]:
    tx(s, 0.8, y, 1.1, 0.5, lab, size=13, bold=True, color=MUTED, anchor="m")
ux = [2.0, 2.65, 3.3, 3.95, 4.6]
for i, x in enumerate(ux):
    box(s, x, 3.95, 0.5, 0.5, str(i + 1), fill=CLIENT_T, line=CLIENT, size=15, bold=True, radius=0.15)
    if i in (1, 3):
        box(s, x, 4.75, 0.5, 0.5, "✕", line="DC2626", dash=True, size=15, bold=True, color="DC2626", radius=0.15)
    else:
        box(s, x, 4.75, 0.5, 0.5, str(i + 1), fill=CLIENT_T if i < 4 else ROOM_T, line=CLIENT if i < 4 else ROOM,
            size=15, bold=True, radius=0.15)
tx(s, 5.15, 4.75, 1.3, 0.5, "← 최신", size=12, bold=True, color=ROOM, anchor="m")
tx(s, 0.8, 5.45, 5.6, 0.8, "최신 5번 위치를 그대로 쓴다.\n지나간 2 · 4는 다시 받지 않는다.", size=13, color="374151")
tx(s, 6.85, 1.62, 5.9, 0.4, "UDP의 약점을 게임이 보완하는 방법", size=15, bold=True)
fixes = [("순서 번호", "번호를 붙이고, 늦게 온 옛날 패킷은 버린다"),
         ("보간", "띄엄띄엄 받은 위치 사이를 부드럽게 이어서 그린다"),
         ("예측", "내 캐릭터는 먼저 움직이고, 서버 값으로 보정한다"),
         ("선택적 재전송", "‘보스 처치’ 같은 중요 이벤트만 확실히 보낸다")]
for i, (k, v) in enumerate(fixes):
    x = 6.85 + (i % 2) * 3.0
    y = 2.1 + (i // 2) * 2.1
    box(s, x, y, 2.88, 1.95, fill=PANEL, line="E5E7EB", radius=0.06)
    tx(s, x + 0.2, y + 0.18, 2.5, 0.45, k, size=18, bold=True, color=ACCENT)
    tx(s, x + 0.2, y + 0.7, 2.5, 1.15, v, size=14)
tx(s, 0.6, 6.45, 12.13, 0.4, "대부분의 게임 네트워크 라이브러리는 UDP 위에 ‘확실한 채널’과 ‘빠른 채널’을 함께 제공한다 → 데이터마다 채널을 고르는 것은 우리 몫",
   size=13, color=MUTED)

# ---------------------------------------------------------------- 13. REST / WebSocket
s = content_slide(P3, "HTTP REST와 WebSocket")
tag(s, "핸드아웃 §3 · §4")
tx(s, 0.6, 1.95, 1.4, 0.9, "응용 방식", size=13, bold=True, color=MUTED, anchor="m")
tx(s, 0.6, 3.05, 1.4, 0.8, "전송 계층", size=13, bold=True, color=MUTED, anchor="m")
box(s, 2.0, 1.95, 1.3, 0.9, "HTTP\nREST", fill=API_T, line=API, size=14, bold=True)
box(s, 3.4, 1.95, 1.4, 0.9, "Web\nSocket", fill=MATCH_T, line=MATCH, size=14, bold=True)
box(s, 4.9, 1.95, 1.45, 0.9, "게임 실시간\n프로토콜", fill=ROOM_T, line=ROOM, size=13, bold=True)
box(s, 2.0, 3.05, 2.8, 0.8, "TCP", fill="334155", size=18, bold=True, color=WHITE)
box(s, 4.9, 3.05, 1.45, 0.8, "UDP", fill="334155", size=18, bold=True, color=WHITE)
tx(s, 0.6, 4.05, 5.8, 1.1, [
    [("REST", {"bold": True, "color": API}), ("로 매칭 완료를 알려면 1초마다 “됐어요?” 하고 물어야 한다", {})],
    [("WebSocket", {"bold": True, "color": MATCH}), ("이면 서버가 “됐습니다”를 바로 보낸다", {})]], size=13, after=6)
table(s, 6.75, 1.95, [1.45, 2.2, 2.33], [
    ["방식", "목적", "예시"],
    [[("HTTP REST", {"color": API})], "요청 후 응답", "프로필 · 결과 저장"],
    [[("WebSocket", {"color": MATCH})], "지속 연결 · 양방향", "로비 알림 · 채팅 · 준비"],
    [[("UDP 기반", {"color": ROOM})], "짧은 패킷 · 빠른 전달", "이동 · 입력 · 스냅샷"],
], row_h=0.85, size=13)
box(s, 0.6, 5.3, 12.13, 1.45, fill=NAVY, radius=0.06)
tx(s, 0.95, 5.4, 11.5, 0.38, [[("✕  ", {"color": "F87171", "bold": True}), ("“REST와 UDP 중 무엇을 쓸까?”  층이 다른 것을 비교하는 질문", {"color": "94A3B8"})]], size=14)
tx(s, 0.95, 5.78, 11.5, 0.5, [[("○  ", {"color": "4ADE80", "bold": True}), ("“이 데이터에는 어떤 전달 특성이 필요한가?”", {"color": WHITE, "bold": True})]], size=21)
tx(s, 1.4, 6.3, 11.0, 0.38, "정확해야 하나 · 빨라야 하나 · 서버가 먼저 알려 줘야 하나", size=14, color="CBD5E1")

# ---------------------------------------------------------------- 14. Firebase 저장
s = content_slide(P4, "Firebase에는 무엇을 저장할까")
tx(s, 0.6, 1.7, 5.5, 1.1, "게임을 껐다 켜도\n남아야 하는 데이터", size=26, bold=True, spacing=1.05)
for i, t in enumerate(["사용자 ID와 프로필", "인벤토리와 설정", "전적과 보상 지급 기록", "검증된 경기 결과"]):
    y = 3.05 + i * 0.72
    box(s, 0.6, y, 0.42, 0.42, "✓", fill=DB_T, shape=MSO_SHAPE.OVAL, size=13, bold=True, color=DB, margin=0)
    tx(s, 1.2, y - 0.04, 4.8, 0.5, t, size=18, bold=(i == 3), color=WARN if i == 3 else INK, anchor="m")
tx(s, 6.55, 1.7, 6.2, 0.4, "보상 · 전적은 이렇게 저장한다", size=17, bold=True)
box(s, 6.55, 2.45, 1.75, 1.0, "클라이언트", fill=CLIENT_T, line=CLIENT, size=15, bold=True)
box(s, 8.75, 2.45, 1.85, 1.0, [{"runs": "서버 측 검증", "bold": True},
                               {"runs": "Cloud Functions 등", "size": 11, "color": MUTED}],
    fill=API_T, line=API, size=15)
box(s, 11.05, 2.35, 1.65, 1.2, "Firebase", fill=DB_T, line=DB, shape=MSO_SHAPE.CAN, size=14, bold=True)
arrow(s, 8.32, 2.95, 8.72, 2.95, color=API)
tx(s, 7.95, 2.12, 1.2, 0.3, "결과 제출", size=11, bold=True, color=API, align="c")
arrow(s, 10.62, 2.95, 11.02, 2.95, color=API)
tx(s, 10.2, 2.12, 1.25, 0.3, "검증된 값만", size=11, bold=True, color=API, align="c")
path(s, [(7.42, 3.47), (7.42, 3.95), (11.87, 3.95), (11.87, 3.57)], color="DC2626", dash=True)
tx(s, 7.6, 4.0, 4.2, 0.35, "✕ 클라이언트가 보상 값을 직접 쓰기", size=12, bold=True, color="DC2626", align="c")
for i, (k, v) in enumerate([("권한 규칙", "보상 · 전적 필드는 클라이언트가 직접 못 쓰게 보안 규칙으로 막는다"),
                            ("중복 처리 기준", "같은 경기 ID의 결과는 한 번만 반영한다")]):
    x = 6.55 + i * 3.13
    box(s, x, 4.65, 3.02, 2.0, fill=PANEL, line="E5E7EB", radius=0.06)
    tx(s, x + 0.2, 4.8, 2.65, 0.4, k, size=17, bold=True, color=ACCENT)
    tx(s, x + 0.2, 5.3, 2.65, 1.3, v, size=14)
box(s, 0.6, 6.0, 5.6, 0.65, "편리함 = 위험: 클라이언트가 DB에 직접 쓸 수 있다", fill=WARN_T, size=14, bold=True, color=WARN)

# ---------------------------------------------------------------- 15. 매칭 상태도
s = content_slide(P4, "기본 매칭 상태도")
tag(s, "1분 토론", color=WARN, fill=WARN_T, x=11.25, w=1.48)
states = [("Lobby", "로비", API, API_T), ("Waiting", "대기", MATCH, MATCH_T), ("Matched", "매칭됨", MATCH, MATCH_T),
          ("Room", "방 입장", ROOM, ROOM_T), ("Playing", "경기 중", ROOM, ROOM_T), ("Result", "결과", API, API_T)]
sw, sx = 1.55, [0.7, 2.75, 4.8, 6.85, 8.9, 10.95]
SY, SH = 2.65, 0.8
for (en, ko, c, t), x in zip(states, sx):
    box(s, x, SY, sw, SH, [{"runs": en, "bold": True, "size": 16}, {"runs": ko, "size": 11, "color": MUTED}],
        fill=t, line=c, lw=1.5, radius=0.25)
fwd = ["매칭 신청", "상대 탐색 성공", "방 입장", "모두 준비", "경기 종료"]
for i, lab in enumerate(fwd):
    x1 = sx[i] + sw
    x2 = sx[i + 1]
    arrow(s, x1 + 0.03, SY + SH / 2, x2 - 0.03, SY + SH / 2, color=INK, w=1.5)
    tx(s, (x1 + x2) / 2 - 0.8, SY - 0.42, 1.6, 0.35, lab, size=11, bold=True, align="c")
path(s, [(sx[5] + sw / 2, SY), (sx[5] + sw / 2, 1.82), (sx[0] + sw / 2, 1.82), (sx[0] + sw / 2, SY)], color=API, w=1.5)
box(s, 5.3, 1.66, 2.6, 0.32, "결과 저장 → 로비", fill=WHITE, size=11, bold=True, color=API, margin=0)
yb = SY + SH
path(s, [(sx[1] + sw / 2, yb), (sx[1] + sw / 2, yb + 0.42), (1.8, yb + 0.42), (1.8, yb)], color=WARN, w=1.5)
tx(s, 1.85, yb + 0.45, 1.6, 0.3, "취소 · 타임아웃", size=11, bold=True, color=WARN, align="c")
path(s, [(sx[2] + sw / 2, yb), (sx[2] + sw / 2, yb + 1.0), (1.1, yb + 1.0), (1.1, yb)], color=WARN, w=1.5)
tx(s, 2.3, yb + 1.03, 2.6, 0.3, "입장 실패 · 상대 이탈", size=11, bold=True, color=WARN, align="c")
tx(s, 6.85, yb + 0.3, 5.7, 0.8, "정상 흐름보다 되돌아가는 화살표에서\n버그가 더 많이 난다", size=14, bold=True, color=WARN)
box(s, 0.6, 5.0, 3.6, 1.8, fill=NAVY, radius=0.06)
tx(s, 0.85, 5.1, 3.2, 0.4, "화살표마다 정한다", size=16, bold=True, color=WHITE)
tx(s, 0.85, 5.55, 3.2, 1.2, ["①  요청 주체", "②  최종 확정 주체", "③  실패 메시지"], size=15, color="E2E8F0", after=3)
table(s, 4.45, 5.0, [2.45, 1.4, 1.4, 3.03], [
    ["화살표", "요청", "확정", "실패하면"],
    ["Waiting → Matched", "매칭 서버", "매칭 서버", "대기 유지 · “상대를 찾는 중”"],
    ["Waiting → Lobby (취소)", "클라이언트", "서버", "서버가 확정한 뒤 화면 전환"],
    [[("Matched → Lobby", {"color": WARN})], [("?", {"color": WARN, "bold": True})],
     [("?", {"color": WARN, "bold": True})], [("함께 정해 보기", {"color": WARN, "bold": True})]],
], row_h=0.43, head_h=0.42, size=13)

# ---------------------------------------------------------------- 16. 매칭 예외
s = content_slide(P4, "매칭 예외와 테스트")
table(s, 0.6, 1.7, [4.4, 7.73], [
    ["예외 상황", "기대하는 최종 상태 (예시)"],
    ["대기 중 취소", "대기열에서 빠지고 로비로 돌아간다"],
    ["취소와 매칭 확정이 동시에 도착", "서버가 먼저 처리한 쪽만 인정 · 두 사람의 상태가 모순되지 않는다"],
    ["매칭 직후 상대 이탈", "남은 사람은 다시 대기열 또는 로비로 (어느 쪽인지는 기획이 정한다)"],
    ["중복 매칭 요청", "버튼을 두 번 눌러도 대기열에는 한 건만 남는다"],
    ["방 입장 실패와 재시도", "정해진 횟수만 재시도하고, 실패하면 이유와 함께 로비로"],
], row_h=0.66, head_h=0.5, size=15)
box(s, 0.6, 5.55, 12.13, 1.2, fill=NAVY, radius=0.06)
tx(s, 0.95, 5.65, 11.5, 0.5, [[("“어떤 오류가 났나?”", {"color": "94A3B8"}), ("보다  ", {"color": "94A3B8"}),
                              ("“최종 상태가 무엇이어야 하나?”", {"color": WHITE, "bold": True}),
                              ("를 먼저 정한다", {"color": WHITE})]], size=19)
tx(s, 0.95, 6.18, 11.5, 0.45, "테스트 문장 예: Waiting에서 취소와 확정이 동시에 오면, 두 사용자의 상태는 서로 모순되지 않아야 한다",
   size=13, color="CBD5E1")

# ---------------------------------------------------------------- 17. 실시간 최적화
s = content_slide(P4, "실시간 데이터와 최적화")
box(s, 0.6, 1.65, 12.13, 0.7, "필요한 데이터만  ·  필요한 대상에게  ·  필요한 주기로", fill="EFF6FF", size=22, bold=True, color=ACCENT)
kinds = [("입력", "이동 · 공격 · 선택", "플레이어가 한 행동\n→ 서버로 보내 처리한다", CLIENT),
         ("상태", "위치 · 체력 · 점수", "최신 값만 중요\n→ 주기적으로 보내고 옛값은 버린다", API),
         ("이벤트", "피격 · 이모트 · 알림", "한 번 일어나고 끝\n→ 확실히 한 번 전달한다", ROOM)]
for i, (k, ex, how, c) in enumerate(kinds):
    x = 0.6 + i * 4.14
    box(s, x, 2.6, 3.85, 2.2, fill=PANEL, line="E5E7EB", radius=0.05)
    box(s, x, 2.6, 0.1, 2.2, fill=c, shape=MSO_SHAPE.RECTANGLE)
    tx(s, x + 0.35, 2.75, 3.3, 0.5, k, size=22, bold=True, color=c)
    tx(s, x + 0.35, 3.3, 3.3, 0.4, ex, size=15, bold=True)
    tx(s, x + 0.35, 3.8, 3.35, 0.95, how, size=13, color="374151")
tips = [("바뀐 값만 보낸다", "체력이 그대로면 체력은 빼고 보낸다"),
        ("멀리 있는 대상은 덜 자주", "시야 밖 플레이어는 갱신 주기를 늘린다"),
        ("늦은 입장자는 스냅샷부터", "현재 상태 + 그 뒤 이벤트만")]
for i, (k, v) in enumerate(tips):
    x = 0.6 + i * 4.14
    tx(s, x, 5.1, 0.6, 0.5, f"{i + 1}", size=26, bold=True, color=ACCENT)
    tx(s, x + 0.55, 5.12, 3.3, 0.45, k, size=17, bold=True)
    tx(s, x + 0.55, 5.58, 3.3, 0.9, v, size=13, color=MUTED)

# ---------------------------------------------------------------- 18. 데모
s = content_slide(P5, "Firebase 매칭 데모")
demo = [("1", "첫 번째 사용자", "매칭 신청 → waiting"),
        ("2", "두 번째 사용자", "매칭 신청 → match 생성"),
        ("3", "두 화면", "상태 변경을 실시간으로 수신")]
for i, (n, who, what) in enumerate(demo):
    y = 1.75 + i * 1.0
    box(s, 0.6, y, 0.6, 0.6, n, fill=ACCENT, shape=MSO_SHAPE.OVAL, size=18, bold=True, color=WHITE, margin=0)
    tx(s, 1.4, y - 0.05, 3.9, 0.35, who, size=13, color=MUTED)
    tx(s, 1.4, y + 0.27, 3.9, 0.45, what, size=17, bold=True)
tx(s, 0.6, 4.8, 4.8, 0.35, [[("열 파일  ", {"color": MUTED}), ("ai_시연/index.html", {"font": MONO, "bold": True})]], size=13)
tx(s, 0.6, 5.12, 4.8, 0.35, "일반 창 + 시크릿 창 = 서로 다른 두 사용자", size=13, color=MUTED)
for i, (title, lines) in enumerate([("일반 창 · 사용자 A", [("① 매칭 신청", "waiting"), ("③ 실시간 수신", "matched")]),
                                    ("시크릿 창 · 사용자 B", [("② 매칭 신청", "match 생성"), ("③ 실시간 수신", "matched")])]):
    x = 5.65 + i * 3.6
    box(s, x, 1.7, 3.4, 3.55, fill=WHITE, line="9CA3AF", radius=0.04)
    box(s, x, 1.7, 3.4, 0.45, fill="E5E7EB", shape=MSO_SHAPE.RECTANGLE)
    for d in range(3):
        box(s, x + 0.15 + d * 0.22, 1.85, 0.14, 0.14, fill=["F87171", "FBBF24", "34D399"][d], shape=MSO_SHAPE.OVAL)
    tx(s, x + 0.85, 1.72, 2.5, 0.4, title, size=12, bold=True, color="374151", anchor="m")
    for j, (k, v) in enumerate(lines):
        y = 2.4 + j * 1.38
        tx(s, x + 0.25, y, 2.9, 0.35, k, size=13, color=MUTED)
        hl = v == "matched"
        box(s, x + 0.25, y + 0.38, 2.9, 0.7, v, fill=GOOD if hl else MATCH_T, line=None if hl else MATCH,
            size=18, bold=True, color=WHITE if hl else MATCH, font=MONO, radius=0.2)
box(s, 8.65, 5.65, 1.2, 0.95, "Firebase", fill=DB_T, line=DB, shape=MSO_SHAPE.CAN, size=12, bold=True)
arrow(s, 7.35, 5.27, 8.62, 6.05, color=DB, tail=True)
arrow(s, 11.15, 5.27, 9.88, 6.05, color=DB, tail=True)
box(s, 0.6, 5.6, 4.8, 1.15, fill=WARN_T, line="FED7AA", radius=0.06)
warn_icon(s, 0.8, 5.75, 0.38)
tx(s, 1.3, 5.68, 4.0, 1.0, [{"runs": "수업용 예제", "bold": True, "color": WARN, "size": 15},
                            "상태 전환을 보여 주는 코드. 전용 게임 서버나 운영용 매칭 완성본이 아니다."],
   size=12, after=2)

# ---------------------------------------------------------------- 19. AI 초안
s = content_slide(P5, "AI로 구조 초안 검토하기")
tx(s, 0.6, 1.65, 12.1, 1.0, [[("AI에게 구조를 결정하게 하지 말고, ", {}), ("빠진 질문을 찾게 한다.", {"color": ACCENT})]],
   size=26, bold=True)
tx(s, 0.6, 2.75, 5.5, 0.4, "요청에 포함할 5가지", size=16, bold=True)
for i, t in enumerate(["경기 인원과 장르", "서버별 책임", "지속 저장 데이터 · 실시간 데이터", "권한과 예외 상황", "원하는 결과물 형식"]):
    y = 3.25 + i * 0.66
    box(s, 0.6, y, 0.45, 0.45, str(i + 1), fill="EFF6FF", shape=MSO_SHAPE.OVAL, size=13, bold=True, color=ACCENT, margin=0)
    tx(s, 1.25, y - 0.04, 4.8, 0.52, t, size=17, anchor="m")
tx(s, 6.6, 2.75, 6.1, 0.4, "요청 예시", size=16, bold=True)
box(s, 6.6, 3.25, 6.13, 1.95, fill="1E293B", radius=0.04, anchor="t", align="l", margin=0.3,
    paras=[[("4인 협동 디펜스 게임이다.", {})],
           [("서버 책임표 · API 목록 · 매칭 상태 · 테스트 시나리오를 만들고", {})],
           [("네가 가정한 부분은 따로 표시해 줘.", {"color": "FCD34D", "bold": True})]],
    size=15, color="E2E8F0", after=8)
tx(s, 6.6, 5.45, 6.1, 0.8, [[("→ ", {"color": ACCENT, "bold": True}),
                            ("답은 정답이 아니라 ", {}), ("팀 검토의 시작점", {"bold": True}), ("으로 쓴다", {})]], size=15)

# ---------------------------------------------------------------- 20. AI 검증
s = content_slide(P5, "AI 답변 검증")
table(s, 0.6, 1.75, [2.4, 5.2], [
    ["확인할 것", "검증 방법"],
    ["서버 책임", "우리 실제 기획 · 구조와 대조한다"],
    ["기술 기능", "공식 문서와 실제 실행으로 확인한다 (기능 · 요금은 자주 바뀐다)"],
    ["권한 · 보안", "규칙을 만들고, 다른 사용자 데이터를 직접 써 본다"],
    ["예외 처리", "취소 · 이탈 · 재접속을 실제로 테스트한다"],
], row_h=0.85, head_h=0.5, size=15)
box(s, 8.5, 1.75, 4.23, 3.9, fill=NAVY, radius=0.05)
tx(s, 8.8, 1.95, 3.7, 0.35, "특히 조심할 문장", size=13, bold=True, color="FCD34D")
tx(s, 8.8, 2.35, 3.7, 0.7, "“서버가 판정한다”", size=26, bold=True, color=WHITE)
tx(s, 8.8, 3.2, 3.7, 0.7, "→ 실제 권한 주체가 누구인지 다시 묻는다", size=16, bold=True, color="E2E8F0")
tx(s, 8.8, 4.05, 3.7, 1.4, "릴레이 서버를 쓰면서 ‘서버가 판정한다’고 적힌 설계가 의외로 많다. 메시지가 서버를 지나간다고 서버가 판정하는 것은 아니다.",
   size=13, color="94A3B8")
tx(s, 0.6, 6.05, 12.13, 0.6, "TCP · UDP · Firebase · Photon에 대한 AI 설명은 제품 문서와 실제 실행으로 확인한다.", size=15, color=MUTED)

# ---------------------------------------------------------------- 21. 설계 순서
s = content_slide("정리", "설계 순서 6단계")
tag(s, "핸드아웃 §6")
order = [("인원 · 장르 · 지연 민감 행동", "몇 명? 0.1초에 민감한 행동은?"),
         ("데이터를 세 종류로 나눈다", "지속 · 매칭 · 실시간 데이터"),
         ("확정 주체와 저장 위치", "누가 확정하고, 어디에 저장하나?"),
         ("API와 실시간 메시지 분리", "요청·응답 목록 / 룸 메시지 목록"),
         ("예외를 상태도로 검토", "취소 · 이탈 · 재접속 · 중복 · 저장 실패"),
         ("최소 인원으로 실제 테스트", "접속 · 지연 · 이탈을 기기 두 대로라도")]
for i, (k, v) in enumerate(order):
    x = 0.6 + (i % 3) * 4.14
    y = 1.7 + (i // 3) * 2.3
    hl = i == 2
    box(s, x, y, 3.85, 2.1, fill="EFF6FF" if hl else PANEL, line=ACCENT if hl else "E5E7EB", lw=1.5 if hl else 1, radius=0.05)
    tx(s, x + 0.25, y + 0.15, 1, 0.6, f"{i + 1}", size=30, bold=True, color=ACCENT)
    tx(s, x + 0.25, y + 0.8, 3.4, 0.5, k, size=17, bold=True)
    tx(s, x + 0.25, y + 1.3, 3.4, 0.7, v, size=13, color=MUTED)
tx(s, 0.6, 6.4, 12.13, 0.45, [[("도구보다 구조가 먼저.  ", {"bold": True, "color": ACCENT}),
                              ("서버 기술을 고르는 것은 거의 마지막 단계다.", {})]], size=16)

# ---------------------------------------------------------------- 22. 질의응답
s = dark_slide()
tx(s, 0.9, 0.9, 11, 1.0, "질의응답", size=44, bold=True, color=WHITE)
tx(s, 0.9, 1.95, 11, 0.5, "질문할 때 알려 주세요", size=18, color="93C5FD", bold=True)
for i, t in enumerate(["현재 게임과 경기 인원", "저장 · 매칭 · 실시간 · 판정 중 막힌 지점", "사용 중인 엔진 · 서비스와 제약"]):
    y = 2.7 + i * 0.85
    box(s, 0.9, y, 0.6, 0.6, str(i + 1), fill=ACCENT, shape=MSO_SHAPE.OVAL, size=18, bold=True, color=WHITE, margin=0)
    tx(s, 1.75, y, 10, 0.6, t, size=22, color=WHITE, anchor="m")
box(s, 0.9, 5.45, 11.53, 1.2, fill="1E293B", radius=0.06)
tx(s, 1.25, 5.55, 11, 0.4, "오늘의 한 문장", size=13, bold=True, color="FCD34D")
tx(s, 1.25, 5.92, 11, 0.6, "“누가 결과를 확정하는가”를 먼저 정하면, 서버 구조와 통신 방식은 따라온다.", size=20, bold=True, color=WHITE)

# ---------------------------------------------------------------- 노트
md = (HERE / "발표대본.md").read_text(encoding="utf-8")
parts = re.split(r"^## 슬라이드 (\d+)\. ", md, flags=re.M)
notes = {}
for i in range(1, len(parts), 2):
    head, _, body = parts[i + 1].partition("\n")
    m = re.search(r"\(([^)]*)\)\s*$", head)
    title = head[: m.start()].strip() if m else head.strip()
    timing = m.group(1) if m else ""
    body = re.sub(r"\*\*(.+?)\*\*", r"\1", body).strip()
    notes[int(parts[i])] = f"[{title} · {timing}]\n\n{body}"

assert len(prs.slides) == TOTAL, len(prs.slides)
for n, slide in enumerate(prs.slides, start=1):
    slide.notes_slide.notes_text_frame.text = notes.get(n, "")

out = HERE / "발표슬라이드.pptx"
prs.save(out)
print("saved", out, len(prs.slides), "slides; notes for", sorted(notes))
